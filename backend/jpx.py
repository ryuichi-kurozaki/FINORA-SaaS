"""TSE listed securities (JPX monthly list) for Japanese code/name search; auto-refreshed weekly check."""
import asyncio
import io
import logging
import re
import unicodedata
from datetime import datetime, timedelta, timezone

import httpx
from openpyxl import load_workbook

from core import db, now_iso

logger = logging.getLogger("jpx")
PAGE = "https://www.jpx.co.jp/markets/statistics-equities/misc/01.html"
SECTOR17 = {"食品": "consumer", "小売": "consumer", "商社・卸売": "consumer", "医薬品": "healthcare", "情報通信・サービスその他": "technology",
            "電機・精密": "technology", "自動車・輸送機": "automotive", "エネルギー資源": "energy", "電力・ガス": "energy", "不動産": "real_estate",
            "銀行": "finance", "金融（除く銀行）": "finance", "運輸・物流": "industrial", "建設・資材": "industrial", "鉄鋼・非鉄": "industrial",
            "素材・化学": "industrial", "機械": "industrial"}
GOLD = ("ゴールド", "金価格", "純金", "金上場", "金先物", "金2倍", "GOLD")
METAL = ("プラチナ", "白金", "純銀", "銀上場", "貴金属")


def norm(s):
    s = unicodedata.normalize("NFKC", s or "").lower().replace(" ", "").replace("\u3000", "")
    return "".join(chr(ord(c) + 0x60) if "\u3041" <= c <= "\u3096" else c for c in s)


def classify(market, name):
    if "内国株式" in market or market == "PRO Market":
        return "jp_stock"
    if "外国株式" in market:
        return "foreign_stock"
    if market.startswith("ETF"):
        if "FANG" not in name and any(w in name for w in GOLD):
            return "gold"
        return "precious_metal" if any(w in name for w in METAL) else "etf"
    if market.startswith("REIT"):
        return "fund" if "インフラ" in name else "real_estate"
    return "other"


def to_item(d):
    return {"ticker": f"{d['code']}.T", "name": d["name"], "exchange": "TSE", "asset_class": d["asset_class"], "country": "JP",
            "sector": d.get("sector"), "currency": "JPY", "source": "jpx"}


async def sync():
    async with httpx.AsyncClient(timeout=60, follow_redirects=True) as h:
        page = (await h.get(PAGE)).text
        path = re.search(r'href="([^"]*data_j\.xlsx?)"', page).group(1)
        raw = (await h.get(httpx.URL(PAGE).join(path))).content
    rows = list(load_workbook(io.BytesIO(raw), read_only=True).active.iter_rows(values_only=True))
    head = [str(c) for c in rows[0]]
    ix = {k: head.index(k) for k in ("日付", "コード", "銘柄名", "市場・商品区分", "17業種区分")}
    batch, docs = now_iso(), []
    for r in rows[1:]:
        if not r[ix["コード"]]:
            continue
        name = unicodedata.normalize("NFKC", str(r[ix["銘柄名"]]))
        market = str(r[ix["市場・商品区分"]])
        cls = classify(market, name)
        docs.append({"code": str(r[ix["コード"]]).strip(), "name": name, "key": norm(name), "market": market, "asset_class": cls,
                     "sector": "real_estate" if cls == "real_estate" else SECTOR17.get(str(r[ix["17業種区分"]])), "batch": batch})
    if len(docs) < 1000:
        raise RuntimeError(f"JPX list too small: {len(docs)}")
    for d in docs:
        await db.jp_securities.update_one({"code": d["code"]}, {"$set": d}, upsert=True)
    await db.jp_securities.delete_many({"batch": {"$ne": batch}})
    as_of = str(rows[1][ix["日付"]])
    await db.system_state.update_one({"id": "jpx"}, {"$set": {"synced_at": batch, "as_of": as_of, "count": len(docs)}}, upsert=True)
    return {"count": len(docs), "as_of": as_of}


async def search(q, limit=8):
    k, code = norm(q), q.strip().upper()
    if not k:
        return []
    found = await db.jp_securities.find({"$or": [{"code": {"$regex": f"^{re.escape(code)}"}}, {"key": {"$regex": re.escape(k)}}]},
                                        {"_id": 0}).limit(60).to_list(60)
    found.sort(key=lambda d: (d["code"] != code, not d["key"].startswith(k), "プライム" not in d["market"], d["code"]))
    return [to_item(d) for d in found[:limit]]


async def get(code):
    d = await db.jp_securities.find_one({"code": code}, {"_id": 0})
    return to_item(d) if d else None


async def loop():
    await db.jp_securities.create_index("code", unique=True)
    await asyncio.sleep(15)
    while True:
        try:
            s = await db.system_state.find_one({"id": "jpx"}) or {}
            if not s.get("synced_at") or datetime.fromisoformat(s["synced_at"]) < datetime.now(timezone.utc) - timedelta(days=7):
                logger.info("JPX list synced: %s", await sync())
        except Exception:
            logger.exception("JPX list sync failed")
        await asyncio.sleep(6 * 3600)
